import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    prs.slide_width = 12192000   # 13.333 inches (16:9)
    prs.slide_height = 6858000  # 7.500 inches

    # Colors
    ORANGE = RGBColor(0xE2, 0x5B, 0x1A)
    ORANGE_LIGHT = RGBColor(0xFA, 0xE8, 0xD8)
    ORANGE_DARK = RGBColor(0xA7, 0x38, 0x07)
    DARK_BG = RGBColor(0x1A, 0x1A, 0x18)
    LIGHT_BG = RGBColor(0xFA, 0xFA, 0xF8)
    CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)
    SURFACE_BG = RGBColor(0xF3, 0xF2, 0xEE)
    TEXT_PRIMARY = RGBColor(0x1A, 0x1A, 0x18)
    TEXT_SECONDARY = RGBColor(0x5A, 0x5A, 0x55)
    TEXT_MUTED = RGBColor(0x8A, 0x8A, 0x84)
    TEXT_LIGHT = RGBColor(0xFF, 0xFF, 0xFF)
    BORDER_COLOR = RGBColor(0xDD, 0xDC, 0xD6)
    GREEN = RGBColor(0x16, 0xA3, 0x4A)
    GREEN_LIGHT = RGBColor(0xDC, 0xFC, 0xE7)
    RED = RGBColor(0xDC, 0x26, 0x26)
    RED_LIGHT = RGBColor(0xFE, 0xE2, 0xE2)
    BLUE = RGBColor(0x18, 0x5F, 0xA5)
    BLUE_LIGHT = RGBColor(0xE6, 0xF1, 0xFB)

    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title, category=None, subtitle=None):
        if category:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
            tf_cat = cat_box.text_frame
            tf_cat.word_wrap = True
            tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category.upper()
            p_cat.font.name = "Segoe UI"
            p_cat.font.size = Pt(11)
            p_cat.font.bold = True
            p_cat.font.color.rgb = ORANGE

        t_top = Inches(0.72) if category else Inches(0.5)
        title_box = slide.shapes.add_textbox(Inches(0.8), t_top, Inches(11.7), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_PRIMARY

        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), t_top + Inches(0.55), Inches(11.7), Inches(0.35))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.name = "Segoe UI"
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_SECONDARY

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def draw_entity_table(slide, x, y, width, table_name, columns, subtitle=None):
        header_h = Inches(0.4)
        row_h = Inches(0.24)
        total_h = header_h + (row_h * len(columns)) + Inches(0.1)

        # Card container
        add_card(slide, x, y, width, total_h, CARD_BG, BORDER_COLOR)

        # Header background banner
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, header_h)
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = ORANGE
        header_shape.line.fill.background()

        tf_h = header_shape.text_frame
        tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_h.margin_left = Inches(0.12)
        p_h = tf_h.paragraphs[0]
        p_h.text = table_name
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_LIGHT

        if subtitle:
            p_sub = tf_h.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Segoe UI"
            p_sub.font.size = Pt(8.5)
            p_sub.font.color.rgb = ORANGE_LIGHT

        # Rows
        cur_y = y + header_h + Inches(0.04)
        for key_type, col_name, col_type in columns:
            tb = slide.shapes.add_textbox(x + Inches(0.08), cur_y, width - Inches(0.16), row_h)
            tf = tb.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            
            # Key badge
            if key_type:
                r_key = p.add_run()
                r_key.text = f"{key_type:<4} "
                r_key.font.name = "Consolas"
                r_key.font.size = Pt(8.5)
                r_key.font.bold = True
                if "PK" in key_type:
                    r_key.font.color.rgb = ORANGE_DARK
                elif "FK" in key_type:
                    r_key.font.color.rgb = BLUE
                else:
                    r_key.font.color.rgb = TEXT_MUTED
            else:
                r_space = p.add_run()
                r_space.text = "     "
                r_space.font.name = "Consolas"
                r_space.font.size = Pt(8.5)

            # Col name
            r_name = p.add_run()
            r_name.text = f"{col_name} "
            r_name.font.name = "Segoe UI"
            r_name.font.size = Pt(9.5)
            r_name.font.bold = True if key_type else False
            r_name.font.color.rgb = TEXT_PRIMARY

            # Col type
            r_type = p.add_run()
            r_type.text = f"({col_type})"
            r_type.font.name = "Consolas"
            r_type.font.size = Pt(8.5)
            r_type.font.color.rgb = TEXT_MUTED

            cur_y += row_h

        return total_h

    # ==========================================
    # SLIDE 1: PORTADA
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, DARK_BG)

    # Accent decorative top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(1.5), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()

    # Project pill
    pill = add_card(s1, Inches(0.8), Inches(1.3), Inches(2.2), Inches(0.4), ORANGE, None)
    tf_pill = pill.text_frame
    tf_pill.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_pill = tf_pill.paragraphs[0]
    p_pill.text = "MESA CLICK"
    p_pill.alignment = PP_ALIGN.CENTER
    p_pill.font.name = "Segoe UI"
    p_pill.font.size = Pt(13)
    p_pill.font.bold = True
    p_pill.font.color.rgb = TEXT_LIGHT

    # Title box
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.5), Inches(2.4))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p1 = tf_title.paragraphs[0]
    p1.text = "Análisis de Base de Datos"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf_title.add_paragraph()
    p2.text = "Justificación del motor de datos relacional (PostgreSQL) y diseño de la estructura del esquema"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(20)
    p2.font.color.rgb = ORANGE_LIGHT

    p3 = tf_title.add_paragraph()
    p3.text = "De la evaluación inicial en MongoDB hacia una arquitectura relacional transaccional y multi-tenant"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(0xBA, 0xBA, 0xB4)

    # Meta footer card
    card_foot = add_card(s1, Inches(0.8), Inches(4.8), Inches(11.7), Inches(1.8), RGBColor(0x24, 0x24, 0x22), None)
    tb_foot = s1.shapes.add_textbox(Inches(1.1), Inches(4.9), Inches(11.0), Inches(1.6))
    tf_foot = tb_foot.text_frame
    tf_foot.word_wrap = True
    
    pf1 = tf_foot.paragraphs[0]
    pf1.text = "ALUMNO / AUTOR"
    pf1.font.name = "Segoe UI"
    pf1.font.size = Pt(10)
    pf1.font.bold = True
    pf1.font.color.rgb = ORANGE

    pf2 = tf_foot.add_paragraph()
    pf2.text = "Pablo Aguirre — Analista en Sistemas"
    pf2.font.name = "Segoe UI"
    pf2.font.size = Pt(15)
    pf2.font.bold = True
    pf2.font.color.rgb = TEXT_LIGHT

    pf3 = tf_foot.add_paragraph()
    pf3.text = "Instituto Remedios de Escalada de San Martín (IREMS) · Práctica Profesionalizante I · Año 2026"
    pf3.font.name = "Segoe UI"
    pf3.font.size = Pt(12)
    pf3.font.color.rgb = RGBColor(0x9A, 0x9A, 0x94)

    # ==========================================
    # SLIDE 2: QUÉ DATOS MANEJA MESA CLICK
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, LIGHT_BG)
    add_header(s2, "¿Qué datos maneja Mesa CLICK?", "1. Relevamiento de Dominio", "Antes de elegir un motor, analizamos la forma y relaciones críticas de nuestros datos")

    cards_data = [
        ("Tenants & Sucursales", "Negocios gastronómicos independientes, multi-sucursal, configuración de horarios y tokens de cobro.", ORANGE),
        ("Sectores & Mesas QR", "Distribución física del salón (Terraza, Interior), mesas con QR estático único y control de estado.", BLUE),
        ("Carta Digital & Variantes", "Categorías, artículos, precios de lista, fotos, grupos de variantes y disponibilidad en tiempo real (86).", GREEN),
        ("Comensales & Pedidos", "Pedidos colaborativos en la misma mesa por dispositivo, notas, comensal identificado y estados de cocina.", ORANGE_DARK),
        ("Cuentas & Pagos", "Cuentas por mesa con control de reapertura (cuenta_version), auditoría y webhooks de Mercado Pago.", BLUE),
        ("Usuarios & Roles", "Acceso administrativo y operativo (admin, encargado, mozo) autenticado mediante Magic Links efímeros.", RGBColor(0x6D, 0x28, 0xD9)),
    ]

    for i, (title, desc, color) in enumerate(cards_data):
        col = i % 3
        row = i // 3
        cx = Inches(0.8 + col * 4.0)
        cy = Inches(1.8 + row * 2.2)
        add_card(s2, cx, cy, Inches(3.7), Inches(2.0), CARD_BG, BORDER_COLOR)
        
        # Color pill
        cp = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.08))
        cp.fill.solid()
        cp.fill.fore_color.rgb = color
        cp.line.fill.background()

        tb = s2.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), Inches(3.3), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_PRIMARY

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_SECONDARY

    # Summary bar at bottom
    add_card(s2, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9), SURFACE_BG, BORDER_COLOR)
    tb_sum = s2.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.8))
    tf_sum = tb_sum.text_frame
    tf_sum.word_wrap = True
    p_sum = tf_sum.paragraphs[0]
    p_sum.text = "Conclusión del relevamiento: Los datos de Mesa CLICK están intrínsecamente relacionados (Tenant -> Sucursal -> Mesa -> Pedido -> Items -> Pagos). Se requiere consistencia ACID estricta para no duplicar pedidos, mezclar mesas ni perder cobros en caja."
    p_sum.font.name = "Segoe UI"
    p_sum.font.size = Pt(11)
    p_sum.font.bold = True
    p_sum.font.color.rgb = TEXT_PRIMARY

    # ==========================================
    # SLIDE 3: DOS FAMILIAS DE BASES DE DATOS
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, LIGHT_BG)
    add_header(s3, "Dos familias de bases de datos", "2. Paradigmas de Persistencia", "Comparación estructural entre modelos Relacionales (SQL) y Orientados a Documentos (NoSQL)")

    # Box SQL
    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), CARD_BG, BORDER_COLOR)
    h_sql = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(0.6))
    h_sql.fill.solid()
    h_sql.fill.fore_color.rgb = ORANGE
    h_sql.line.fill.background()
    tf_hsql = h_sql.text_frame
    tf_hsql.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_hsql.margin_left = Inches(0.3)
    p_hsql = tf_hsql.paragraphs[0]
    p_hsql.text = "SQL Relacional (PostgreSQL / MySQL)"
    p_hsql.font.name = "Segoe UI"
    p_hsql.font.size = Pt(15)
    p_hsql.font.bold = True
    p_hsql.font.color.rgb = TEXT_LIGHT

    tb_sql_body = s3.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(5.0), Inches(3.8))
    tf_sql_body = tb_sql_body.text_frame
    tf_sql_body.word_wrap = True
    sql_points = [
        ("Tablas, filas y columnas:", "Estructura rigurosa y normalizada definida de antemano."),
        ("Claves foráneas (FK):", "Integridad referencial que impide ítems huérfanos o pedidos sin mesa."),
        ("Transacciones ACID nativas:", "Atomicidad y aislamiento total: cobros y cambios de estado 'todo o nada'."),
        ("Consultas y Agregaciones:", "SQL declarativo optimizado mediante JOINs, índices B-Tree y métricas GROUP BY."),
        ("Migraciones formales:", "Evolución controlada y trazable mediante archivos SQL versionados.")
    ]
    for i, (title, desc) in enumerate(sql_points):
        p = tf_sql_body.paragraphs[0] if i == 0 else tf_sql_body.add_paragraph()
        r1 = p.add_run()
        r1.text = title + " "
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = TEXT_PRIMARY
        r2 = p.add_run()
        r2.text = desc + "\n"
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_SECONDARY

    # Box NoSQL
    add_card(s3, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), CARD_BG, BORDER_COLOR)
    h_nosql = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(0.6))
    h_nosql.fill.solid()
    h_nosql.fill.fore_color.rgb = RGBColor(0x3B, 0x82, 0xF6)
    h_nosql.line.fill.background()
    tf_hnosql = h_nosql.text_frame
    tf_hnosql.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_hnosql.margin_left = Inches(0.3)
    p_hnosql = tf_hnosql.paragraphs[0]
    p_hnosql.text = "NoSQL de Documentos (MongoDB / Firestore)"
    p_hnosql.font.name = "Segoe UI"
    p_hnosql.font.size = Pt(15)
    p_hnosql.font.bold = True
    p_hnosql.font.color.rgb = TEXT_LIGHT

    tb_nosql_body = s3.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(5.0), Inches(3.8))
    tf_nosql_body = tb_nosql_body.text_frame
    tf_nosql_body.word_wrap = True
    nosql_points = [
        ("Documentos BSON / JSON:", "Colecciones donde cada documento puede tener campos y tipos distintos."),
        ("Datos anidados (Embedding):", "Agrupa ítems dentro del pedido en lugar de relacionar tablas."),
        ("Sin claves foráneas estrictas:", "La consistencia de datos recae completamente en el código de backend."),
        ("Transacciones distribuidas:", "Soporta transacciones multi-documento pero con mayor costo y complejidad."),
        ("Escalabilidad horizontal nativa:", "Diseñado para volúmenes masivos de datos desnormalizados e independientes.")
    ]
    for i, (title, desc) in enumerate(nosql_points):
        p = tf_nosql_body.paragraphs[0] if i == 0 else tf_nosql_body.add_paragraph()
        r1 = p.add_run()
        r1.text = title + " "
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = TEXT_PRIMARY
        r2 = p.add_run()
        r2.text = desc + "\n"
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 4: CUADRO COMPARATIVO SQL VS NOSQL
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, LIGHT_BG)
    add_header(s4, "Cuadro comparativo: SQL vs NoSQL", "3. Evaluación Técnica", "Análisis riguroso de requerimientos frente a las necesidades reales de Mesa CLICK")

    # Table
    rows = 7
    cols = 3
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(4.8)

    tbl_shape = s4.shapes.add_table(rows, cols, left, top, width, height)
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.8)
    tbl.columns[1].width = Inches(4.45)
    tbl.columns[2].width = Inches(4.45)

    headers = ["Criterio Operativo", "SQL Relacional (PostgreSQL) - ELEGIDO", "NoSQL Documentos (MongoDB) - DESCARTADO"]
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ORANGE if c_idx == 1 else (DARK_BG if c_idx == 0 else RGBColor(0x4A, 0x4A, 0x48))
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = h_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    comp_data = [
        ("Modelo de Datos y Multi-Tenancy", "Encaja perfecto: Tenant -> Sucursal -> Sector -> Mesa -> Pedidos con FK estrictas.", "Obliga a duplicar tenant/sucursal en cada documento o usar $lookup manual."),
        ("Consistencia e Integridad", "Claves foráneas y checks evitan pedidos huérfanos, mesas inválidas o cuentas duplicadas.", "Sin FK: la integridad queda en manos del código, riesgoso en altas de cocina y salón."),
        ("Concurrencia en Mesa (Múltiples Celulares)", "Aislamiento transaccional y bloqueo por fila: varios comensales agregan ítems sin colisión.", "Actualizar arrays concurrentes en un mismo documento genera colisiones y 'lost updates'."),
        ("Snapshots Históricos de Precios", "Separa artículo del precio histórico unitario en pedido_items de forma inmutable.", "Requiere desnormalizar el artículo completo dentro de la comanda para no perder histórico."),
        ("Reportes, Métricas & Facturación (S17)", "SQL estándar con GROUP BY, SUM(), índices B-Tree ultra rápidos y analítica nativa.", "Pipelines de agregación ($unwind, $group) verbosos, difíciles de mantener y pesados."),
        ("Flexibilidad Semiestructurada", "Soporte híbrido con JSONB nativo para horarios de turnos y webhooks de Mercado Pago.", "Esquema flexible, pero la heterogeneidad de documentos complica el código Go en el back.")
    ]

    for r_idx, row_vals in enumerate(comp_data):
        for c_idx, val in enumerate(row_vals):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if (r_idx % 2 == 0) else SURFACE_BG
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.12)
            tf.margin_top = Inches(0.06)
            p = tf.paragraphs[0]
            p.text = val
            p.font.name = "Segoe UI"
            p.font.size = Pt(9.5)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = TEXT_PRIMARY
            elif c_idx == 1:
                p.font.color.rgb = RGBColor(0x0F, 0x68, 0x2A)
            else:
                p.font.color.rgb = RGBColor(0x8A, 0x24, 0x24)

    # ==========================================
    # SLIDE 5: MATRIZ DE PONDERACIÓN
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, LIGHT_BG)
    add_header(s5, "Motores evaluados: Matriz de ponderación", "4. Selección de Tecnología", "Puntaje cuantitativo de 1 a 5 según requerimientos técnicos de Mesa CLICK")

    # Table
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(11.7)
    height = Inches(3.6)

    tbl_shape5 = s5.shapes.add_table(7, 5, left, top, width, height)
    tbl5 = tbl_shape5.table
    tbl5.columns[0].width = Inches(3.7)
    tbl5.columns[1].width = Inches(2.0)
    tbl5.columns[2].width = Inches(2.0)
    tbl5.columns[3].width = Inches(2.0)
    tbl5.columns[4].width = Inches(2.0)

    heads5 = ["Criterio de Evaluación (1 al 5)", "PostgreSQL", "MongoDB", "MySQL / MariaDB", "SQLite"]
    for c_idx, h in enumerate(heads5):
        cell = tbl5.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ORANGE if c_idx == 1 else DARK_BG
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]
        p.text = h
        p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    scores = [
        ("Ajuste a datos relacionales y Multi-Tenancy", "5", "2", "5", "4"),
        ("Integridad referencial y transacciones ACID", "5", "3", "5", "3"),
        ("Soporte para JSON semiestructurado nativo", "5 (JSONB)", "5", "4", "3"),
        ("Concurrencia, Pool y performance en Go", "5 (pgx nativo)", "3", "4", "2"),
        ("Métricas de negocio y agregaciones (KPIs)", "5", "2", "4", "2"),
        ("TOTAL (sobre 25 puntos)", "25 / 25", "15 / 25", "22 / 25", "14 / 25")
    ]

    for r_idx, s_row in enumerate(scores):
        for c_idx, val in enumerate(s_row):
            cell = tbl5.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            is_total = (r_idx == len(scores) - 1)
            if is_total:
                cell.fill.fore_color.rgb = ORANGE_LIGHT if c_idx == 1 else SURFACE_BG
            else:
                cell.fill.fore_color.rgb = CARD_BG if (r_idx % 2 == 0) else SURFACE_BG
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.1)
            p = tf.paragraphs[0]
            p.text = val
            p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT
            p.font.name = "Segoe UI"
            p.font.size = Pt(11 if is_total else 10)
            p.font.bold = True if (is_total or c_idx == 1) else False
            if c_idx == 1:
                p.font.color.rgb = ORANGE_DARK if is_total else ORANGE
            else:
                p.font.color.rgb = TEXT_PRIMARY

    # Bottom analysis card
    add_card(s5, Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.4), CARD_BG, BORDER_COLOR)
    tb_diag = s5.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.2))
    tf_diag = tb_diag.text_frame
    tf_diag.word_wrap = True
    p_d1 = tf_diag.paragraphs[0]
    p_d1.text = "Lectura del resultado:"
    p_d1.font.name = "Segoe UI"
    p_d1.font.size = Pt(11)
    p_d1.font.bold = True
    p_d1.font.color.rgb = ORANGE

    p_d2 = tf_diag.add_paragraph()
    p_d2.text = "• PostgreSQL resulta el ganador indiscutido (25/25): Ofrece el rigor relacional necesario para la facturación y los pedidos en vivo, sumado al poder de JSONB para configuraciones flexibles.\n• MongoDB (15/25): Descartado porque obligaba a implementar la integridad por software, penalizando la concurrencia de comensales y complicando las consultas analíticas del negocio.\n• MySQL (22/25): Excelente alternativa relacional, pero PostgreSQL ofrece tipos avanzados (UUID nativo, JSONB indexable GIN) y mejor integración concurrente con el stack Go."
    p_d2.font.name = "Segoe UI"
    p_d2.font.size = Pt(10)
    p_d2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 6: POR QUÉ SE PENSÓ EN MONGO DB
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, LIGHT_BG)
    add_header(s6, "¿Por qué en su momento se pensó en MongoDB?", "5. El Dilema Inicial", "La hipótesis inicial del equipo y la aparente afinidad con el dominio gastronómico")

    reasons = [
        ("1. La ilusión del 'Menú como un JSON'", "En una primera mirada, una carta gastronómica parece un árbol de documentos:\n\n• Una categoría contiene platos.\n• Un plato contiene ingredientes, variantes y opciones.\n• Parecía intuitivo guardar todo el menú de un restaurante como un único documento JSON estructurado.", ORANGE),
        ("2. El pedido como comanda anidada", "El flujo de un cliente al escanear la mesa parece un envío de formulario:\n\n• El cliente selecciona varios ítems con notas.\n• Parecía natural hacer un solo `insert` de un documento comanda con un array `items: [{...}]` sin preocuparse por normalizar ni hacer múltiples inserciones relacionales.", BLUE),
        ("3. Velocidad inicial sin migraciones (Schemaless)", "La promesa de desarrollo ágil inicial:\n\n• Sin necesidad de definir DDL o tipos estrictos al comienzo.\n• Facilidad para agregar campos sobre la marcha sin ejecutar comandos `ALTER TABLE`.\n• Familiaridad en la industria para MVPs sencillos basados en JavaScript/Node.js.", GREEN)
    ]

    for i, (title, desc, color) in enumerate(reasons):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.8)
        add_card(s6, cx, cy, Inches(3.7), Inches(4.0), CARD_BG, BORDER_COLOR)

        top_line = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.08))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = color
        top_line.line.fill.background()

        tb = s6.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.25), Inches(3.3), Inches(3.6))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_SECONDARY

    # Warning alert box at bottom
    add_card(s6, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.0), RED_LIGHT, RGBColor(0xFC, 0xA5, 0xA5))
    tb_alert = s6.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.9))
    tf_alert = tb_alert.text_frame
    tf_alert.word_wrap = True
    p_al = tf_alert.paragraphs[0]
    p_al.text = "EL QUIEBRE DE LA HIPÓTESIS:"
    p_al.font.name = "Segoe UI"
    p_al.font.size = Pt(10.5)
    p_al.font.bold = True
    p_al.font.color.rgb = RED

    p_al2 = tf_alert.add_paragraph()
    p_al2.text = "Lo que parecía simple para mostrar una carta estática colapsó al modelar la dinámica real de un salón: pedidos colaborativos concurrentes en la misma mesa, stock (86), auditoría de cambios de precios y liquidación de pagos. MongoDB transfería toda la responsabilidad de consistencia a la capa de aplicación."
    p_al2.font.name = "Segoe UI"
    p_al2.font.size = Pt(10)
    p_al2.font.color.rgb = TEXT_PRIMARY

    # ==========================================
    # SLIDE 7: POR QUÉ SE DESCARTÓ MONGO DB
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, LIGHT_BG)
    add_header(s7, "¿Por qué se descartó MongoDB?", "6. Análisis Crítico", "Los cuatro problemas estructurales que hicieron inviable MongoDB para Mesa CLICK")

    pitfalls = [
        ("1. Concurrencia en Mesa y 'Lost Updates'", "En Mesa CLICK varios comensales piden a la vez desde sus celulares al mismo QR de mesa. Si guardamos el pedido como un único documento con un array de ítems, dos comensales enviando órdenes concurrentes provocan condiciones de carrera y sobrescritura de platos."),
        ("2. El dilema de la actualización de precios", "Si el local cambia el precio de un plato en el catálogo: ¿Cómo garantizamos que los pedidos abiertos o pasados no cambien su total? En Mongo obligaría a clonar manualmente todos los datos del producto dentro del pedido; en SQL basta una columna inmutable `precio_unitario`."),
        ("3. Falta de Foreign Keys e Integridad Huérfana", "Si un encargado elimina una categoría o desactiva una mesa con comensales comiendo, MongoDB no previene referencias rotas. La integridad referencial dependería 100% de validar por código, multiplicando el riesgo de bugs en producción."),
        ("4. Agregaciones y Métricas de Negocio Complejas", "Calcular facturación diaria por sucursal, ticket promedio y ranking de artículos (Sprint 17) en Mongo exige pipelines `$lookup` y `$unwind` muy pesados. En PostgreSQL son simples sentencias `SUM()`, `JOIN` y `GROUP BY` resueltas en milisegundos.")
    ]

    for i, (title, desc) in enumerate(pitfalls):
        col = i % 2
        row = i // 2
        cx = Inches(0.8 + col * 6.0)
        cy = Inches(1.8 + row * 2.5)
        add_card(s7, cx, cy, Inches(5.7), Inches(2.3), CARD_BG, BORDER_COLOR)

        # Red badge on card
        rb = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(0.12), Inches(2.3))
        rb.fill.solid()
        rb.fill.fore_color.rgb = RED
        rb.line.fill.background()

        tb = s7.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.15), Inches(5.2), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 8: POR QUÉ POSTGRESQL ES IDEAL
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, LIGHT_BG)
    add_header(s8, "¿Por qué PostgreSQL es el motor ideal?", "7. Justificación de la Elección", "Robustez relacional, transaccionalidad ACID y soporte híbrido semiestructurado")

    reasons_pg = [
        ("Transacciones ACID Estrictas", "Garantía total de consistencia: registrar un pedido, asociar sus ítems y notificar por SSE se ejecuta de forma atómica.", GREEN),
        ("Soporte Híbrido: JSONB Nativo", "Almacena horarios de sucursales y payloads crudos de webhooks de Mercado Pago con indexación binaria sin perder la estructura SQL.", ORANGE),
        ("Integración Nativa con Go", "Driver `pgx` de alto rendimiento con pool de conexiones ligero, perfectamente integrado al modelo concurrente de Goroutines.", BLUE),
        ("Tipos Nativos y Seguridad", "Uso de `UUID` nativo (`gen_random_uuid()`), `TIMESTAMPTZ` para manejo preciso de husos horarios y constraints `CHECK`.", RGBColor(0x6D, 0x28, 0xD9)),
        ("Migraciones SQL Versionadas", "Control de versiones exacto del esquema (`001_...` a `021_...`), reproducible en desarrollo, staging y producción en Render.", ORANGE_DARK),
        ("Cero Costo de Licencia", "Motor 100% de código abierto, sin vendor lock-in, estándar mundial soportado por cualquier nube y plataforma.", GREEN)
    ]

    for i, (title, desc, color) in enumerate(reasons_pg):
        col = i % 3
        row = i // 3
        cx = Inches(0.8 + col * 4.0)
        cy = Inches(1.8 + row * 2.4)
        add_card(s8, cx, cy, Inches(3.7), Inches(2.2), CARD_BG, BORDER_COLOR)

        top_b = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.08))
        top_b.fill.solid()
        top_b.fill.fore_color.rgb = color
        top_b.line.fill.background()

        tb = s8.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), Inches(3.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 9: DECISIÓN OFICIAL Y STACK
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, LIGHT_BG)
    add_header(s9, "Decisión Final y Stack Tecnológico", "8. Arquitectura Adoptada", "La sinergia entre Go, PostgreSQL y Next.js para un SaaS gastronómico en tiempo real")

    # Big 3 cards
    layers = [
        ("Base de Datos", "PostgreSQL 16+", [
            "Motor relacional desplegado en Render.",
            "Esquema estructurado en 14 tablas principales.",
            "Campos JSONB estratégicos para flexibilidad.",
            "Migraciones SQL puras sin ORM opaco.",
            "Consultas parametrizadas seguras contra SQLi."
        ], ORANGE),
        ("Backend Core", "Go (Golang) + net/http", [
            "Alta concurrencia con Goroutines ligeras.",
            "Logging estructurado nativo con `slog`.",
            "Canales Server-Sent Events (SSE) en vivo.",
            "Autenticación stateless JWT + Magic Links.",
            "Pool de conexiones optimizado con `pgx`."
        ], BLUE),
        ("Frontend PWA", "Next.js 16 + Tailwind CSS", [
            "Experiencia Mobile-First para comensales.",
            "Dashboard reactivo para mozos y cocina.",
            "Comunicación en tiempo real vía SSE.",
            "Rutas dinámicas por mesa mediante código QR.",
            "Diseño sin fricción (cero instalación requerida)."
        ], GREEN)
    ]

    for i, (layer, tech, bullets, color) in enumerate(layers):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.8)
        add_card(s9, cx, cy, Inches(3.7), Inches(4.8), CARD_BG, BORDER_COLOR)

        h_sh = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.7))
        h_sh.fill.solid()
        h_sh.fill.fore_color.rgb = color
        h_sh.line.fill.background()
        tf_h = h_sh.text_frame
        tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_h.margin_left = Inches(0.2)
        p1 = tf_h.paragraphs[0]
        p1.text = layer.upper()
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(0xFF, 0xEE, 0xDD)

        p2 = tf_h.add_paragraph()
        p2.text = tech
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT

        tb = s9.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.9), Inches(3.3), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        for b_idx, b in enumerate(bullets):
            p = tf.paragraphs[0] if b_idx == 0 else tf.add_paragraph()
            p.text = "• " + b
            p.font.name = "Segoe UI"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 10: MODELO ENTIDAD-RELACIÓN GENERAL
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, LIGHT_BG)
    add_header(s10, "Modelo Entidad-Relación: Visión General", "9. Diseño de Datos", "Estructura normalizada organizada en 3 subsistemas funcionales independientes pero acoplados")

    # Big metrics card
    add_card(s10, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.1), SURFACE_BG, BORDER_COLOR)
    metrics = [
        ("14 Tablas Relacionales", "Esquema normalizado y modular"),
        ("21 Migraciones SQL", "Control de versiones atómico"),
        ("18 Foreign Keys con CASCADE", "Integridad referencial total"),
        ("Columnas Híbridas JSONB", "Flexibilidad en horarios y pagos")
    ]
    for i, (m_title, m_desc) in enumerate(metrics):
        mx = Inches(1.0 + i * 2.9)
        tb_m = s10.shapes.add_textbox(mx, Inches(1.9), Inches(2.7), Inches(0.9))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        pm1 = tf_m.paragraphs[0]
        pm1.text = m_title
        pm1.font.name = "Segoe UI"
        pm1.font.size = Pt(13)
        pm1.font.bold = True
        pm1.font.color.rgb = ORANGE

        pm2 = tf_m.add_paragraph()
        pm2.text = m_desc
        pm2.font.name = "Segoe UI"
        pm2.font.size = Pt(9.5)
        pm2.font.color.rgb = TEXT_SECONDARY

    # 3 Subsystems cards
    subs = [
        ("Módulo 1: Organización & Auth", "Administración del negocio multi-tenant y control de usuarios.", [
            "• `tenants`: Negocio raíz, marca, redes (JSONB).",
            "• `usuarios`: Personal con roles (admin, mozo).",
            "• `magic_tokens`: Tokens de login de 15 min.",
            "• `sucursales`: Puntos de venta y horarios (JSONB)."
        ], ORANGE),
        ("Módulo 2: Salón & Carta Digital", "Configuración física del salón y catálogo gastronómico.", [
            "• `sectores` y `mesas`: QR estático y control de cuenta.",
            "• `categorias`: Agrupación y orden visual de la carta.",
            "• `articulos`: Platos, precios, fotos y estado activo.",
            "• `variantes`: Grupos, selección única y adicionales."
        ], BLUE),
        ("Módulo 3: Operación & Pagos", "Ciclo de vida de la comanda, comensales y cobro online.", [
            "• `pedidos`: Estados (recibido -> preparando -> listo).",
            "• `pedido_items`: Snapshot de precio unitario y notas.",
            "• `comensales`: Identificación por dispositivo en mesa.",
            "• `pagos`: Integración y conciliación Mercado Pago."
        ], GREEN)
    ]

    for i, (title, subtitle, bullets, color) in enumerate(subs):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(3.1)
        add_card(s10, cx, cy, Inches(3.7), Inches(3.5), CARD_BG, BORDER_COLOR)

        top_l = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.55))
        top_l.fill.solid()
        top_l.fill.fore_color.rgb = color
        top_l.line.fill.background()
        tf_tl = top_l.text_frame
        tf_tl.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tl.margin_left = Inches(0.2)
        pt = tf_tl.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT

        tb = s10.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.65), Inches(3.3), Inches(2.7))
        tf = tb.text_frame
        tf.word_wrap = True
        ps = tf.paragraphs[0]
        ps.text = subtitle + "\n"
        ps.font.name = "Segoe UI"
        ps.font.size = Pt(9.5)
        ps.font.italic = True
        ps.font.color.rgb = TEXT_MUTED

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10)
            pb.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 11: DER MÓDULO 1 - ORGANIZACIÓN Y AUTH
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11, LIGHT_BG)
    add_header(s11, "DER Módulo 1: Organización, Multi-Tenancy y Acceso", "10. Esquema Detallado", "Aislamiento por negocio, gestión de sucursales y autenticación sin contraseña (Magic Link)")

    draw_entity_table(s11, Inches(0.8), Inches(1.8), Inches(2.7), "tenants", [
        ("PK", "id", "UUID"),
        ("", "nombre", "TEXT"),
        ("", "nombre_fantasia", "TEXT"),
        ("", "rubro", "TEXT"),
        ("", "slug", "TEXT UQ"),
        ("", "redes_sociales", "JSONB"),
        ("", "mp_access_token", "TEXT"),
        ("", "created_at", "TIMESTAMPTZ")
    ], "Entidad raíz del negocio")

    draw_entity_table(s11, Inches(3.8), Inches(1.8), Inches(2.7), "usuarios", [
        ("PK", "id", "UUID"),
        ("FK", "tenant_id", "UUID"),
        ("UQ", "email", "TEXT"),
        ("", "nombre", "TEXT"),
        ("", "rol", "TEXT (admin|mozo)"),
        ("", "created_at", "TIMESTAMPTZ")
    ], "Personal interno del local")

    draw_entity_table(s11, Inches(6.8), Inches(1.8), Inches(2.5), "magic_tokens", [
        ("PK", "id", "UUID"),
        ("FK", "usuario_id", "UUID"),
        ("UQ", "token", "TEXT"),
        ("", "expires_at", "TIMESTAMPTZ"),
        ("", "used_at", "TIMESTAMPTZ")
    ], "Login passwordless (15 min)")

    draw_entity_table(s11, Inches(9.6), Inches(1.8), Inches(2.8), "sucursales", [
        ("PK", "id", "UUID"),
        ("FK", "tenant_id", "UUID"),
        ("", "nombre", "TEXT"),
        ("", "whatsapp", "TEXT"),
        ("", "telefono", "TEXT"),
        ("", "horarios", "JSONB"),
        ("", "mp_access_token", "TEXT"),
        ("", "created_at", "TIMESTAMPTZ")
    ], "Puntos de venta físicos")

    # Bottom notes card
    add_card(s11, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.3), SURFACE_BG, BORDER_COLOR)
    tb_n1 = s11.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.1))
    tf_n1 = tb_n1.text_frame
    tf_n1.word_wrap = True
    pn1 = tf_n1.paragraphs[0]
    pn1.text = "Decisiones Clave de Diseño:"
    pn1.font.name = "Segoe UI"
    pn1.font.size = Pt(11)
    pn1.font.bold = True
    pn1.font.color.rgb = ORANGE

    pn2 = tf_n1.add_paragraph()
    pn2.text = "• Multi-tenancy lógico mediante `tenant_id`: Todas las consultas operativas se filtran por tenant, impidiendo fugas de datos entre distintos locales.\n• Magic Tokens de un solo uso (`used_at`): Elimina contraseñas en texto plano o hashes vulnerables. El token expira automáticamente en 15 minutos.\n• Horarios en `JSONB`: Permite almacenar múltiples turnos por día de la semana sin necesidad de crear una tabla rígida de horarios."
    pn2.font.name = "Segoe UI"
    pn2.font.size = Pt(9.5)
    pn2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 12: DER MÓDULO 2 - SALÓN Y CARTA
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, LIGHT_BG)
    add_header(s12, "DER Módulo 2: Salón y Carta Gastronómica", "11. Esquema Detallado", "Mesas identificadas por token QR y catálogo de platos con variantes y disponibilidad")

    draw_entity_table(s12, Inches(0.8), Inches(1.8), Inches(2.1), "sectores", [
        ("PK", "id", "UUID"),
        ("FK", "sucursal_id", "UUID"),
        ("", "nombre", "TEXT")
    ], "Zonas (Salón, Terraza)")

    draw_entity_table(s12, Inches(3.1), Inches(1.8), Inches(2.6), "mesas", [
        ("PK", "id", "UUID"),
        ("FK", "sucursal_id", "UUID"),
        ("FK", "sector_id", "UUID"),
        ("", "numero", "INT"),
        ("", "capacidad", "INT"),
        ("UQ", "qr_token", "TEXT"),
        ("", "estado", "TEXT"),
        ("", "cuenta_version", "INT"),
        ("", "pago_habilitado", "BOOLEAN")
    ], "Mesas físicas con QR")

    draw_entity_table(s12, Inches(5.9), Inches(1.8), Inches(2.2), "categorias", [
        ("PK", "id", "UUID"),
        ("FK", "tenant_id", "UUID"),
        ("", "nombre", "TEXT"),
        ("", "orden", "INT")
    ], "Entradas, Bebidas, etc.")

    draw_entity_table(s12, Inches(8.3), Inches(1.8), Inches(2.4), "articulos", [
        ("PK", "id", "UUID"),
        ("FK", "tenant_id", "UUID"),
        ("FK", "categoria_id", "UUID"),
        ("", "nombre", "TEXT"),
        ("", "descripcion", "TEXT"),
        ("", "precio", "NUMERIC(10,2)"),
        ("", "foto_url", "TEXT"),
        ("", "activo", "BOOLEAN")
    ], "Platos y productos")

    draw_entity_table(s12, Inches(10.9), Inches(1.8), Inches(2.0), "variantes", [
        ("PK", "id", "UUID"),
        ("FK", "articulo_id", "UUID"),
        ("", "grupo", "TEXT"),
        ("", "nombre", "TEXT"),
        ("", "precio_adic.", "NUMERIC"),
        ("", "seleccion_un.", "BOOLEAN"),
        ("", "orden", "INT")
    ], "Opciones y extras")

    # Bottom notes card
    add_card(s12, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.3), SURFACE_BG, BORDER_COLOR)
    tb_n2 = s12.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.1))
    tf_n2 = tb_n2.text_frame
    tf_n2.word_wrap = True
    pn1 = tf_n2.paragraphs[0]
    pn1.text = "Decisiones Clave de Diseño:"
    pn1.font.name = "Segoe UI"
    pn1.font.size = Pt(11)
    pn1.font.bold = True
    pn1.font.color.rgb = ORANGE

    pn2 = tf_n2.add_paragraph()
    pn2.text = "• `qr_token` único e inmutable: La mesa física mantiene un código QR impreso permanente. No se expone el ID interno en la URL pública.\n• `cuenta_version`: Permite agrupar los pedidos de los comensales que están sentados actualmente, independizándolos de consumos anteriores.\n• Grupos de variantes (`grupo`, `seleccion_unica`): Soporta tanto opciones excluyentes (punto de carne) como agregados múltiples (aderezos extra)."
    pn2.font.name = "Segoe UI"
    pn2.font.size = Pt(9.5)
    pn2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 13: DER MÓDULO 3 - PEDIDOS Y PAGOS
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s13, LIGHT_BG)
    add_header(s13, "DER Módulo 3: Pedidos, Comensales y Pagos", "12. Esquema Detallado", "Gestión de pedidos colaborativos, snapshot inmutable de precios e integración con Mercado Pago")

    draw_entity_table(s13, Inches(0.8), Inches(1.8), Inches(2.7), "pedidos", [
        ("PK", "id", "UUID"),
        ("FK", "mesa_id", "UUID"),
        ("FK", "sucursal_id", "UUID"),
        ("", "cuenta_version", "INT"),
        ("", "estado", "TEXT (recibido|listo)"),
        ("", "created_at", "TIMESTAMPTZ"),
        ("", "updated_at", "TIMESTAMPTZ")
    ], "Comanda global de cocina")

    draw_entity_table(s13, Inches(3.7), Inches(1.8), Inches(3.0), "pedido_items", [
        ("PK", "id", "UUID"),
        ("FK", "pedido_id", "UUID"),
        ("FK", "articulo_id", "UUID"),
        ("", "comensal_id", "UUID"),
        ("", "comensal_nombre", "VARCHAR(100)"),
        ("", "cantidad", "INT"),
        ("", "precio_unitario", "NUMERIC(10,2)"),
        ("", "notas", "TEXT")
    ], "Platos pedidos con snapshot")

    draw_entity_table(s13, Inches(6.9), Inches(1.8), Inches(2.4), "pedido_item_variantes", [
        ("PK", "id", "UUID"),
        ("FK", "pedido_item_id", "UUID"),
        ("FK", "variante_id", "UUID")
    ], "Variantes elegidas")

    draw_entity_table(s13, Inches(9.5), Inches(1.8), Inches(3.0), "pagos", [
        ("PK", "id", "UUID"),
        ("FK", "mesa_id", "UUID"),
        ("", "cuenta_version", "INT"),
        ("", "proveedor", "VARCHAR(50)"),
        ("", "preferencia_id", "VARCHAR(100)"),
        ("", "pago_id", "VARCHAR(100)"),
        ("", "monto", "NUMERIC(10,2)"),
        ("", "estado", "VARCHAR(50)"),
        ("", "detalles", "JSONB"),
        ("", "created_at", "TIMESTAMPTZ")
    ], "Historial y webhooks MP")

    # Bottom notes card
    add_card(s13, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.3), SURFACE_BG, BORDER_COLOR)
    tb_n3 = s13.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.1))
    tf_n3 = tb_n3.text_frame
    tf_n3.word_wrap = True
    pn1 = tf_n3.paragraphs[0]
    pn1.text = "Decisiones Clave de Diseño:"
    pn1.font.name = "Segoe UI"
    pn1.font.size = Pt(11)
    pn1.font.bold = True
    pn1.font.color.rgb = ORANGE

    pn2 = tf_n3.add_paragraph()
    pn2.text = "• Snapshot en `precio_unitario`: El precio se copia de `articulos` al momento exacto de crear el pedido. Si el admin cambia el precio mañana, el balance histórico no se corrompe jamás.\n• Soporte colaborativo (`comensal_id`, `comensal_nombre`): Permite que varios clientes agreguen platos a la comanda general y luego dividan la cuenta por persona.\n• Tabla `pagos` con `cuenta_version`: Asegura que el pago registrado en Mercado Pago salde exactamente los pedidos de esa versión de cuenta."
    pn2.font.name = "Segoe UI"
    pn2.font.size = Pt(9.5)
    pn2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 14: DESAFÍO TÉCNICO 1 - CONCURRENCIA Y CUENTAS
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s14, LIGHT_BG)
    add_header(s14, "Desafío Técnico 1: Concurrencia y Cuentas en Mesa", "13. Casos de Uso Críticos", "¿Cómo resolver pedidos simultáneos y rotación de clientes en la misma mesa física?")

    cards_dt1 = [
        ("El Problema del Salón", "En un restaurante real ocurren dos fenómenos críticos:\n\n1. Concurrencia masiva: 4 amigos escanean el mismo QR al mismo tiempo y envían platos a la cocina simultáneamente.\n2. Rotación continua: Cuando la mesa 5 paga y se va, llega otro grupo de clientes y escanea exactamente el mismo código QR impreso en la mesa.", RED),
        ("La Falla de Enfoques Tradicionales", "• En MongoDB: Guardar los pedidos dentro del documento de la mesa generaría 'lost updates' cuando dos comensales confirman a la vez.\n• Borrar pedidos al liberar: Se perdería todo el registro contable y la facturación histórica del día.\n• Cambiar el QR en cada cliente: Imposible operativamente en gastronomía (los QR están plastificados o grabados en madera).", ORANGE_DARK),
        ("Nuestra Solución: 'cuenta_version'", "Introdujimos la columna `cuenta_version INTEGER` en `mesas`, `pedidos` y `pagos`:\n\n• Mientras la mesa está abierta, todos los comensales aportan a `cuenta_version = N`.\n• Al solicitar la cuenta y cerrar la mesa, el sistema incrementa `cuenta_version = N + 1`.\n• Los pedidos anteriores quedan cerrados e históricos; los nuevos comensales empiezan una cuenta en blanco con el mismo QR.", GREEN)
    ]

    for i, (title, desc, color) in enumerate(cards_dt1):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.8)
        add_card(s14, cx, cy, Inches(3.7), Inches(4.8), CARD_BG, BORDER_COLOR)

        top_bar = s14.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.6))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = color
        top_bar.line.fill.background()
        tf_tb = top_bar.text_frame
        tf_tb.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tb.margin_left = Inches(0.2)
        pt = tf_tb.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT

        tb = s14.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.8), Inches(3.3), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        pd = tf.paragraphs[0]
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(11)
        pd.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 15: DESAFÍO TÉCNICO 2 - FOTOS DE PLATOS
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s15, LIGHT_BG)
    add_header(s15, "Desafío Técnico 2: Fotos de la Carta (¿BD o Cloud?)", "14. Arquitectura de Archivos", "Estimación de impacto del almacenamiento multimedia y decisión de arquitectura")

    # Left: Numbers & Comparison
    add_card(s15, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), CARD_BG, BORDER_COLOR)
    tb_l = s15.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    
    p = tf_l.paragraphs[0]
    p.text = "Estimación para 50 locales (carta de 100 platos c/u):"
    p.font.name = "Segoe UI"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_PRIMARY

    p2 = tf_l.add_paragraph()
    p2.text = "• 5.000 imágenes de platos en alta resolución (promedio ~2.5 MB c/u).\n• Volumen total multimedia: ≈ 12.5 Gigabytes de imágenes."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_SECONDARY

    # 2 Comparison mini cards inside
    p3 = tf_l.add_paragraph()
    p3.text = "\nOpción A: Guardar binario en PostgreSQL (BYTEA)"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = RED

    p4 = tf_l.add_paragraph()
    p4.text = "✘ La base pesa 12.5 GB en disco.\n✘ Los backups tardan horas y saturan la RAM del servidor Go en cada consulta.\n✘ Cada SELECT a la carta transfiere megabytes innecesarios."
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(10)
    p4.font.color.rgb = TEXT_SECONDARY

    p5 = tf_l.add_paragraph()
    p5.text = "\nOpción B: Guardar solo URL en PostgreSQL (ELEGIDA)"
    p5.font.name = "Segoe UI"
    p5.font.size = Pt(11)
    p5.font.bold = True
    p5.font.color.rgb = GREEN

    p6 = tf_l.add_paragraph()
    p6.text = "✔ La base pesa menos de 2 MB (5.000 filas × ~200 bytes de texto).\n✔ Backups instantáneos, consultas ultra rápidas.\n✔ Imágenes servidas a comensales por CDN optimizado para móviles."
    p6.font.name = "Segoe UI"
    p6.font.size = Pt(10)
    p6.font.color.rgb = TEXT_SECONDARY

    # Right: Architecture diagram card
    add_card(s15, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), CARD_BG, BORDER_COLOR)
    tb_r = s15.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    pr1 = tf_r.paragraphs[0]
    pr1.text = "Flujo de Carga y Visualización:"
    pr1.font.name = "Segoe UI"
    pr1.font.size = Pt(13)
    pr1.font.bold = True
    pr1.font.color.rgb = ORANGE

    flow_steps = [
        ("1. Carga desde el Admin (Next.js)", "El encargado sube la foto del plato desde el dashboard web."),
        ("2. Almacenamiento en Nube / CDN", "La imagen se procesa, optimiza y aloja en almacenamiento de objetos (Cloudinary / S3 / R2)."),
        ("3. Persistencia en PostgreSQL", "En la tabla `articulos`, solo se guarda el string en la columna `foto_url TEXT`."),
        ("4. Consumo Mobile Comensal", "El celular del cliente descarga la imagen directamente desde el CDN perimetral, sin sobrecargar la API en Go ni la base de datos.")
    ]
    for st_title, st_desc in flow_steps:
        p_st = tf_r.add_paragraph()
        r1 = p_st.add_run()
        r1.text = "\n" + st_title + "\n"
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_PRIMARY
        r2 = p_st.add_run()
        r2.text = st_desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 16: MODELO HÍBRIDO TABLAS VS JSONB
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s16, LIGHT_BG)
    add_header(s16, "Enfoque Híbrido: ¿Cuándo usar Tablas y cuándo JSONB?", "15. Pragmatismo de Diseño", "La ventaja clave de PostgreSQL: no tener que elegir ciegamente entre SQL y NoSQL")

    cases = [
        ("Tablas Normalizadas (SQL Puro)", "Para entidades core y datos transaccionales", [
            "• `pedidos` y `pedido_items`",
            "• `articulos` y `categorias`",
            "• `mesas` y `usuarios`",
            "",
            "¿Por qué SQL puro?",
            "Requieren unicidad, integridad referencial (FK), agregaciones matemáticas (`SUM`, `COUNT`) y control de concurrencia ACID."
        ], BLUE),
        ("Columnas Semiestructuradas (JSONB)", "Para datos variables o integración con terceros", [
            "• `sucursales.horarios`:",
            "  `{ lunes: '09:00-23:00', ... }` (flexibilidad sin tablas intermedias de turnos).",
            "• `tenants.redes_sociales`:",
            "  `{ ig: '@bar', web: '...' }`.",
            "• `pagos.detalles`:",
            "  Payload crudo de respuesta de Mercado Pago para auditoría técnica."
        ], ORANGE),
        ("Lo mejor de dos mundos", "Elimina la necesidad de MongoDB", [
            "• PostgreSQL almacena JSONB en formato binario descompuesto.",
            "• Permite indexar claves internas de JSON mediante índices GIN.",
            "• Permite hacer consultas que cruzan tablas SQL tradicionales con filtros sobre campos dentro del JSON.",
            "• Evita la complejidad de mantener dos motores de BD distintos."
        ], GREEN)
    ]

    for i, (title, subtitle, bullets, color) in enumerate(cases):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.8)
        add_card(s16, cx, cy, Inches(3.7), Inches(4.8), CARD_BG, BORDER_COLOR)

        top_b = s16.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.6))
        top_b.fill.solid()
        top_b.fill.fore_color.rgb = color
        top_b.line.fill.background()
        tf_tb = top_b.text_frame
        tf_tb.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tb.margin_left = Inches(0.2)
        pt = tf_tb.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT

        tb = s16.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.75), Inches(3.3), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        ps = tf.paragraphs[0]
        ps.text = subtitle + "\n"
        ps.font.name = "Segoe UI"
        ps.font.size = Pt(10)
        ps.font.italic = True
        ps.font.color.rgb = TEXT_MUTED

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = TEXT_PRIMARY if "•" not in b and b != "" else TEXT_SECONDARY

    # ==========================================
    # SLIDE 17: ESTRATEGIA DE MIGRACIONES Y CI/CD
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s17, LIGHT_BG)
    add_header(s17, "Estrategia de Migraciones y Despliegue", "16. DevOps y Evolución", "Cómo evoluciona el esquema de forma segura y determinista en cada despliegue a producción")

    mig_cards = [
        ("1. Migraciones SQL Puras Versionadas", "En `repos/api/migrations/` residen los archivos numerados secuencialmente:\n\n• `001_create_tenants.sql`\n• `006_create_mesas.sql`\n• `010_create_pedidos.sql`\n• `016_add_comensal_pedidos.sql`\n• `019_create_pagos.sql` ... hasta `021`.\n\nSin magia de ORMs que generen esquemas impredecibles.", ORANGE),
        ("2. Ejecución Atómica en el Startup", "El servidor Go corre las migraciones pendientes automáticamente al iniciar:\n\n• Se ejecuta dentro de una transacción `BEGIN ... COMMIT`.\n• Si una migración falla, el servidor aborta el inicio sin dejar la base en estado corrupto.\n• Registro de versiones aplicadas en tabla de control interno.", BLUE),
        ("3. Despliegue Continuo en Render", "Integración completa con el Git Workflow:\n\n• Flujo de ramas: `feat/*` -> `qa` -> `main`.\n• Al mergear a `main`, el Dockerfile compila el backend en Go.\n• El contenedor en Render se conecta a PostgreSQL, aplica nuevas migraciones y levanta el servicio sin tiempo de inactividad.", GREEN)
    ]

    for i, (title, desc, color) in enumerate(mig_cards):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(1.8)
        add_card(s17, cx, cy, Inches(3.7), Inches(4.8), CARD_BG, BORDER_COLOR)

        top_b = s17.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.55))
        top_b.fill.solid()
        top_b.fill.fore_color.rgb = color
        top_b.line.fill.background()
        tf_tb = top_b.text_frame
        tf_tb.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tb.margin_left = Inches(0.2)
        pt = tf_tb.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_LIGHT

        tb = s17.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.7), Inches(3.3), Inches(3.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pd = tf.paragraphs[0]
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 18: EVOLUCIÓN FUTURA Y ANALÍTICA (S14-S19)
    # ==========================================
    s18 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s18, LIGHT_BG)
    add_header(s18, "Preparación para el Futuro: Analítica y KDS", "17. Roadmap del Esquema", "Cómo la base de datos actual soporta la evolución de la Fase 4 de Mesa CLICK")

    sprints_data = [
        ("Sprint 14: Disponibilidad (86) & Horarios", "Gestión de ítems no disponibles (`activo = false`) y cartas segmentadas por franja horaria combinando `sucursales.horarios` y estado de apertura."),
        ("Sprint 15: KDS Cocina & Comandas", "Pantalla de cocina interactiva en tiempo real; los cambios de estado por ítem y mesa se despachan vía SSE sincronizados con PostgreSQL."),
        ("Sprint 16: Modelo Freemium (Free vs Pro)", "Validación de cuotas de suscripción en middleware (límite de mesas, sucursales y artículos) mediante conteos SQL instantáneos."),
        ("Sprint 17-18: Dashboard de Métricas y Analítica", "Consultas analíticas optimizadas (facturación, ticket promedio, horas pico, platos más vendidos) aprovechando índices compuestos en PostgreSQL.")
    ]

    for i, (title, desc) in enumerate(sprints_data):
        col = i % 2
        row = i // 2
        cx = Inches(0.8 + col * 6.0)
        cy = Inches(1.8 + row * 2.4)
        add_card(s18, cx, cy, Inches(5.7), Inches(2.2), CARD_BG, BORDER_COLOR)

        sb = s18.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(0.12), Inches(2.2))
        sb.fill.solid()
        sb.fill.fore_color.rgb = ORANGE
        sb.line.fill.background()

        tb = s18.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.15), Inches(5.2), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_PRIMARY

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_SECONDARY

    # ==========================================
    # SLIDE 19: CONCLUSIÓN Y CIERRE
    # ==========================================
    s19 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s19, DARK_BG)

    # Decorative bar
    bar19 = s19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(1.5), Inches(0.08))
    bar19.fill.solid()
    bar19.fill.fore_color.rgb = ORANGE
    bar19.line.fill.background()

    # Title
    tb_c = s19.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.5), Inches(1.2))
    tf_c = tb_c.text_frame
    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "Conclusiones del Análisis de Base de Datos"
    p_c1.font.name = "Segoe UI"
    p_c1.font.size = Pt(32)
    p_c1.font.bold = True
    p_c1.font.color.rgb = TEXT_LIGHT

    # 3 Summary Pillar Cards
    pillars = [
        ("BASE DE DATOS", "SQL Relacional: PostgreSQL", "Nuestros datos son intrínsecamente relacionales y de estricta consistencia. Descartar MongoDB evitó inconsistencias contables y 'lost updates' en mesas concurrentes. PostgreSQL garantiza ACID, transacciones robustas y soporte JSONB de primera clase.", ORANGE),
        ("MULTIMEDIA", "Archivos en la Nube + CDN", "Las imágenes de los platos nunca se guardan como binarios en la BD. Se alojan en almacenamiento de objetos en la nube y se sirven por CDN, manteniendo la base de datos ultra liviana (< 2 MB) con backups instantáneos.", BLUE),
        ("ARQUITECTURA", "Go + Migraciones SQL Puras", "Cero dependencias opacas. Control total del DDL mediante 21 migraciones versionadas que corren de forma determinista en el startup. El sistema escala eficientemente en producción sobre Render.", GREEN)
    ]

    for i, (tag, title, desc, color) in enumerate(pillars):
        cx = Inches(0.8 + i * 4.0)
        cy = Inches(2.6)
        add_card(s19, cx, cy, Inches(3.7), Inches(3.6), RGBColor(0x25, 0x25, 0x22), None)

        top_l = s19.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, Inches(3.7), Inches(0.08))
        top_l.fill.solid()
        top_l.fill.fore_color.rgb = color
        top_l.line.fill.background()

        tb = s19.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), Inches(3.3), Inches(3.2))
        tf = tb.text_frame
        tf.word_wrap = True

        ptag = tf.paragraphs[0]
        ptag.text = tag
        ptag.font.name = "Segoe UI"
        ptag.font.size = Pt(9.5)
        ptag.font.bold = True
        ptag.font.color.rgb = color

        ptit = tf.add_paragraph()
        ptit.text = title
        ptit.font.name = "Segoe UI"
        ptit.font.size = Pt(13)
        ptit.font.bold = True
        ptit.font.color.rgb = TEXT_LIGHT

        pdesc = tf.add_paragraph()
        pdesc.text = desc
        pdesc.font.name = "Segoe UI"
        pdesc.font.size = Pt(10)
        pdesc.font.color.rgb = RGBColor(0xBA, 0xBA, 0xB4)

    # Footer note
    tb_fn = s19.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4))
    tf_fn = tb_fn.text_frame
    p_fn = tf_fn.paragraphs[0]
    p_fn.text = "Mesa CLICK · Repositorio: github.com/aguirrepablo-iresm/mesa-click · IREMS 2026"
    p_fn.font.name = "Segoe UI"
    p_fn.font.size = Pt(10)
    p_fn.font.color.rgb = RGBColor(0x7A, 0x7A, 0x74)

    output_path = "Analisis_BD_Mesa_Click.pptx"
    prs.save(output_path)
    print(f"Presentation successfully generated at: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
